"""Actualiza secciones concretas de la memoria sin reconstruir el documento completo."""

from __future__ import annotations

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "docs/memoria/TFM_definitivo.docx"


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def _set_run_font(run, name: str, size: float, color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_inline(paragraph, text: str) -> None:
    token_pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)")
    cursor = 0
    for match in token_pattern.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, "Consolas", 9.5)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def _move_before(element, boundary) -> None:
    boundary.addprevious(element)


def _add_paragraph_before(doc, boundary, text: str = "", style: str = "normal"):
    paragraph = doc.add_paragraph(style=style)
    if text:
        _add_inline(paragraph, text)
    _move_before(paragraph._p, boundary)
    return paragraph


def _add_code_before(doc, boundary, lines: list[str]) -> None:
    paragraph = doc.add_paragraph(style="normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        _set_run_font(run, "Consolas", 8.7)
        if index < len(lines) - 1:
            run.add_break()
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    properties.append(shading)
    _move_before(paragraph._p, boundary)


def _add_table_before(doc, boundary, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=max(len(row) for row in rows))
    table.autofit = True
    _set_table_borders(table)
    _repeat_header(table.rows[0])
    for row_index, values in enumerate(rows):
        _prevent_row_split(table.rows[row_index])
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            if row_index == 0:
                _set_cell_shading(cell, "1F4E78")
            elif row_index % 2 == 0:
                _set_cell_shading(cell, "EAF2F8")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if column_index > 0 else WD_ALIGN_PARAGRAPH.LEFT
            )
            _add_inline(paragraph, value)
            for run in paragraph.runs:
                _set_run_font(run, "Times New Roman", 8.3, "FFFFFF" if row_index == 0 else "000000")
                if row_index == 0:
                    run.bold = True
    _move_before(table._tbl, boundary)
    spacer = _add_paragraph_before(doc, boundary)
    spacer.paragraph_format.space_after = Pt(2)


def _add_image_before(doc, boundary, image_path: Path, caption: str | None) -> None:
    paragraph = doc.add_paragraph(style="normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = bool(caption)
    paragraph.add_run().add_picture(str(image_path), width=Inches(6.1))
    _move_before(paragraph._p, boundary)
    if caption:
        caption_paragraph = _add_paragraph_before(doc, boundary, caption)
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.paragraph_format.keep_together = True
        for run in caption_paragraph.runs:
            run.italic = True
            run.font.size = Pt(10)


def _is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _insert_markdown(doc, boundary, markdown_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    index = 1 if lines and lines[0].startswith("# ") else 0
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        text = " ".join(part.strip() for part in paragraph_buffer).strip()
        paragraph_buffer.clear()
        paragraph = _add_paragraph_before(doc, boundary, text)
        paragraph.paragraph_format.space_after = Pt(6)

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("```"):
            flush_paragraph()
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            _add_code_before(doc, boundary, code_lines)
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and _is_table_separator(lines[index + 1]):
            flush_paragraph()
            table_lines = [stripped]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [
                [cell.strip() for cell in table_line.strip("|").split("|")]
                for table_line in table_lines
            ]
            _add_table_before(doc, boundary, rows)
            continue
        image_match = re.fullmatch(r"!\[(.+?)\]\((.+?)\)", stripped)
        if image_match:
            flush_paragraph()
            caption = image_match.group(1) if image_match.group(1).startswith("Figura") else None
            image_path = (markdown_path.parent / image_match.group(2)).resolve()
            _add_image_before(doc, boundary, image_path, caption)
            index += 1
            continue
        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            flush_paragraph()
            level = len(heading_match.group(1)) + 1
            paragraph = _add_paragraph_before(doc, boundary, heading_match.group(2), f"Heading {min(level, 4)}")
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                _set_run_font(
                    run,
                    "Times New Roman",
                    run.font.size.pt if run.font.size else 12,
                    "000000",
                )
            index += 1
            continue
        list_match = re.match(r"^([-*]|\d+\.)\s+(.+)$", stripped)
        if list_match:
            flush_paragraph()
            marker, item = list_match.groups()
            continuation: list[str] = [item]
            index += 1
            while index < len(lines):
                candidate = lines[index].strip()
                if not candidate or re.match(r"^([-*]|\d+\.)\s+", candidate) or candidate.startswith(("#", "|", "```", "![")):
                    break
                continuation.append(candidate)
                index += 1
            prefix = "• " if marker in {"-", "*"} else f"{marker} "
            paragraph = _add_paragraph_before(doc, boundary, prefix + " ".join(continuation))
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.paragraph_format.space_after = Pt(3)
            continue
        paragraph_buffer.append(stripped)
        index += 1
    flush_paragraph()


def _find_paragraph(doc, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"No se encontró el encabezado: {prefix}")


def _polish_existing_layout(doc) -> None:
    """Evita desbordamientos en elementos conservados del apartado 4.2."""
    for table in doc.tables:
        if not table.rows or table.cell(0, 0).text.strip() != "Grupo":
            continue
        for row in table.rows:
            _prevent_row_split(row)
            for cell in row.cells:
                _set_cell_margins(cell, top=55, start=70, bottom=55, end=70)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        _set_run_font(
                            run,
                            "Arial",
                            8.5,
                            "FFFFFF" if row is table.rows[0] else "000000",
                        )

    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("La evaluación mantiene el orden cronológico"):
            paragraph.paragraph_format.keep_together = True


def _replace_between(doc, start_prefix: str, end_prefix: str, markdown_path: Path) -> None:
    start = _find_paragraph(doc, start_prefix)
    end = _find_paragraph(doc, end_prefix)
    body = doc._element.body
    children = list(body)
    start_index = children.index(start._p)
    end_index = children.index(end._p)
    for element in children[start_index + 1 : end_index]:
        body.remove(element)
    _insert_markdown(doc, end._p, markdown_path)


def main() -> None:
    document = Document(DOCX_PATH)
    _replace_between(
        document,
        "3.1.",
        "3.2.",
        ROOT / "docs/memoria/03_01_fuentes_variables_calidad_datos.md",
    )
    _replace_between(
        document,
        "3.2.",
        "3.3.",
        ROOT / "docs/memoria/03_02_arquitectura_tecnologias_flujo.md",
    )
    _replace_between(
        document,
        "3.3.",
        "4. Procesamiento",
        ROOT / "docs/memoria/03_03_planificacion_organizacion_proyecto.md",
    )
    _replace_between(
        document,
        "4.1.",
        "4.2.",
        ROOT / "docs/memoria/04_01_extraccion_transformacion_almacenamiento.md",
    )
    _replace_between(
        document,
        "4.3.",
        "5. Sistema",
        ROOT / "docs/memoria/04_03_desarrollo_evaluacion_modelo_predictivo.md",
    )
    _polish_existing_layout(document)

    temporary = DOCX_PATH.with_name("TFM_definitivo.actualizado.docx")
    document.save(temporary)
    verification = Document(temporary)
    if not any(p.text.startswith("3.1.4.") for p in verification.paragraphs):
        raise RuntimeError("La actualización del apartado 3.1 no quedó completa.")
    if not any(p.text.startswith("3.2.4.") for p in verification.paragraphs):
        raise RuntimeError("La actualización del apartado 3.2 no quedó completa.")
    if not any("pimiento al aire libre" in p.text.lower() for p in verification.paragraphs):
        raise RuntimeError("No se encontró la definición del piloto al aire libre.")
    if not any(p.text.startswith("4.1.11.") for p in verification.paragraphs):
        raise RuntimeError("La actualización del apartado 4.1 no quedó completa.")
    if not any("XGBoost" in p.text for p in verification.paragraphs):
        raise RuntimeError("La actualización del apartado 4.3 no quedó completa.")
    os.replace(temporary, DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
